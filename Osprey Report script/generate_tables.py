import tkinter as tk
from tkinter import filedialog, messagebox
import pandas as pd
import openpyxl
import os
import glob
from openpyxl.utils.dataframe import dataframe_to_rows

def get_tier(metric, value):
    if pd.isna(value):
        return ""
    if metric == "Thermal Efficiency w/o char":
        if value >= 50: return "Tier 5"
        elif value >= 40: return "Tier 4"
        elif value >= 30: return "Tier 3"
        elif value >= 20: return "Tier 2"
        elif value >= 10: return "Tier 1"
        else: return "Tier 0"
    elif metric == "Thermal Efficiency w char":
        if value >= 50: return "Tier 5"
        elif value >= 40: return "Tier 4"
        elif value >= 30: return "Tier 3"
        elif value >= 20: return "Tier 2"
        elif value >= 10: return "Tier 1"
        else: return "Tier 0"
    elif metric == "CO per useful energy":
        if value <= 3.0: return "Tier 5"
        elif value <= 4.4: return "Tier 4"
        elif value <= 7.2: return "Tier 3"
        elif value <= 11.5: return "Tier 2"
        elif value <= 18.3: return "Tier 1"
        else: return "Tier 0"
    elif metric == "PM2.5 per useful energy":
        if value <= 5: return "Tier 5"
        elif value <= 62: return "Tier 4"
        elif value <= 218: return "Tier 3"
        elif value <= 481: return "Tier 2"
        elif value <= 1030: return "Tier 1"
        else: return "Tier 0"
    return ""

def process_tests(root_dir, selected_folders):
    # Mapping of "Table Variable Name" -> CSV Variable Names (hp, mp, lp, weighted)
    var_map = {
        "Thermal Efficiency w/o char": {"High": "eff_wo_char_hp", "Medium": "eff_wo_char_mp", "Low": "eff_wo_char_lp", "Average": "eff_wo_char_weighted"},
        "Thermal Efficiency w char": {"High": "eff_w_char_hp", "Medium": "eff_w_char_mp", "Low": "eff_w_char_lp", "Average": "eff_w_char_weighted"},
        "Fire Power": {"High": "firepower_w_char_hp", "Medium": "firepower_w_char_mp", "Low": "firepower_w_char_lp", "Average": "cooking_power_weighted"},
        "PM2.5 per useful energy": {"High": "PM_useful_eng_deliver_hp", "Medium": "PM_useful_eng_deliver_mp", "Low": "PM_useful_eng_deliver_lp", "Average": "PM_useful_eng_deliver_weighted"},
        "PM2.5 mass per time": {"High": "PM_mass_time_hp", "Medium": "PM_mass_time_mp", "Low": "PM_mass_time_lp", "Average": "PM_mass_time_weighted"},
        "CO per useful energy": {"High": "CO_useful_eng_deliver_hp", "Medium": "CO_useful_eng_deliver_mp", "Low": "CO_useful_eng_deliver_lp", "Average": "CO_useful_eng_deliver_weighted"},
        "CO mass per time": {"High": "CO_mass_time_hp", "Medium": "CO_mass_time_mp", "Low": "CO_mass_time_lp", "Average": "CO_mass_time_weighted"},
        "Time to boil": {"High": "time_to_boil_hp", "Medium": "time_to_boil_mp", "Low": "time_to_boil_lp", "Average": None}
    }
    
    # Store data
    extracted_data = {t: {} for t in selected_folders}
    
    # Extract data
    for folder in selected_folders:
        csv_pattern = os.path.join(folder, "*AllOutputs.csv")
        csv_files = glob.glob(csv_pattern)
        if not csv_files:
            csv_pattern = os.path.join(folder, "*allOutputs.csv")
            csv_files = glob.glob(csv_pattern)
        
        if not csv_files:
            messagebox.showerror("Error", f"No AllOutputs.csv found in {folder}")
            return
        
        df = pd.read_csv(csv_files[0])
        for index, row in df.iterrows():
            extracted_data[folder][row['variable_name']] = row['value']
            
    # Calculate Averages
    avg_data = {}
    for metric, phases in var_map.items():
        avg_data[metric] = {}
        for phase, csv_var in phases.items():
            if not csv_var:
                avg_data[metric][phase] = float('nan')
                continue
                
            vals = []
            for folder in selected_folders:
                val = extracted_data[folder].get(csv_var)
                try:
                    vals.append(float(val))
                except:
                    pass
            if vals:
                avg_data[metric][phase] = sum(vals) / len(vals)
            else:
                avg_data[metric][phase] = float('nan')

    # Update Excel
    template_path = os.path.join(os.path.dirname(__file__), "table templates.xlsx")
    wb = openpyxl.load_workbook(template_path)
    
    # Update Series Avg
    ws_avg = wb["Series Avg"]
    
    # Fill in the number of replicates in cell B1
    ws_avg.cell(row=1, column=2).value = len(selected_folders)
    
    for row in range(3, 12):
        metric = ws_avg.cell(row=row, column=1).value
        if isinstance(metric, str):
            metric = metric.strip()
            ws_avg.cell(row=row, column=1).value = metric
            
        if metric and metric in avg_data:
            ws_avg.cell(row=row, column=3).value = avg_data[metric]["High"]
            ws_avg.cell(row=row, column=4).value = avg_data[metric]["Medium"]
            ws_avg.cell(row=row, column=5).value = avg_data[metric]["Low"]
            ws_avg.cell(row=row, column=6).value = avg_data[metric]["Average"]
            
            # Tier Rating in column G
            tier = get_tier(metric, avg_data[metric]["Average"])
            if tier:
                ws_avg.cell(row=row, column=7).value = tier

    # Update Multi-Test All Tests
    if len(selected_folders) > 1:
        ws_all = wb["Multi-Test All Tests"]
        
        phase_map = {"HIGH POWER": "High", "MEDIUM POWER": "Medium", "LOW POWER": "Low", "TIER RATING": "Average"}
        current_phase = None
        
        for row in range(2, 34):
            metric = ws_all.cell(row=row, column=1).value
            if isinstance(metric, str):
                metric = metric.strip()
                ws_all.cell(row=row, column=1).value = metric
                
            if metric in phase_map:
                current_phase = phase_map[metric]
                continue
                
            # Mapping TIER RATING metric names
            if current_phase == "Average":
                tier_metric_map = {
                    "Efficiency without char": "Thermal Efficiency w/o char",
                    "Efficiency with char": "Thermal Efficiency w char",
                    "PM useful energy delivered": "PM2.5 per useful energy",
                    "CO useful energy delivered": "CO per useful energy"
                }
                real_metric = tier_metric_map.get(metric)
            else:
                real_metric = metric
                
            if real_metric and real_metric in avg_data:
                for i, folder in enumerate(selected_folders):
                    csv_var = var_map[real_metric][current_phase]
                    if csv_var:
                        val = extracted_data[folder].get(csv_var)
                        try:
                            val = float(val)
                            if current_phase == "Average":
                                val = get_tier(real_metric, val)
                            ws_all.cell(row=row, column=3+i).value = val
                        except:
                            pass
                
                avg_val = avg_data[real_metric][current_phase]
                if current_phase == "Average":
                    ws_all.cell(row=row, column=8).value = get_tier(real_metric, avg_val)
                else:
                    ws_all.cell(row=row, column=8).value = avg_val
    else:
        # Delete Multi-Test All Tests
        if "Multi-Test All Tests" in wb.sheetnames:
            del wb["Multi-Test All Tests"]

    # Delete other unused sheets
    if "var names" in wb.sheetnames: del wb["var names"]
    if "ISO Tiers" in wb.sheetnames: del wb["ISO Tiers"]

    root_folder_name = os.path.basename(os.path.normpath(root_dir))
    test_type = "multi-test" if len(selected_folders) > 1 else "single-test"
    filename = f"{root_folder_name}_{test_type}_Summary.xlsx"
    save_path = os.path.join(root_dir, filename)
    wb.save(save_path)
    messagebox.showinfo("Success", f"Tables generated successfully at:\n{save_path}")

class App:
    def __init__(self, root):
        self.root = root
        self.root.title("Osprey Table Generator")
        self.root.geometry("400x500")
        
        tk.Label(root, text="Step 1: Select Root Folder", font=("Arial", 12, "bold")).pack(pady=10)
        
        self.btn_browse = tk.Button(root, text="Browse...", command=self.browse_root)
        self.btn_browse.pack()
        
        self.lbl_root = tk.Label(root, text="No folder selected", fg="gray")
        self.lbl_root.pack(pady=5)
        
        tk.Label(root, text="Step 2: Select up to 5 Replicates", font=("Arial", 12, "bold")).pack(pady=10)
        
        self.listbox_frame = tk.Frame(root)
        self.listbox_frame.pack(fill=tk.BOTH, expand=True, padx=20)
        
        self.check_vars = {}
        
        self.btn_generate = tk.Button(root, text="Generate Tables", font=("Arial", 12, "bold"), bg="green", fg="white", command=self.generate)
        self.btn_generate.pack(pady=20)
        self.btn_generate.config(state=tk.DISABLED)
        
    def browse_root(self):
        folder = filedialog.askdirectory(title="Select Root Folder")
        if folder:
            self.lbl_root.config(text=folder, fg="black")
            self.root_folder = folder
            self.load_replicates(folder)
            
    def load_replicates(self, root_folder):
        for widget in self.listbox_frame.winfo_children():
            widget.destroy()
            
        self.check_vars = {}
        subdirs = [f.path for f in os.scandir(root_folder) if f.is_dir()]
        
        valid_dirs = []
        for d in subdirs:
            if glob.glob(os.path.join(d, "*AllOutputs.csv")) or glob.glob(os.path.join(d, "*allOutputs.csv")):
                valid_dirs.append(d)
                
        if not valid_dirs:
            tk.Label(self.listbox_frame, text="No replicates found with AllOutputs.csv", fg="red").pack()
            self.btn_generate.config(state=tk.DISABLED)
            return
            
        for d in valid_dirs:
            var = tk.BooleanVar()
            chk = tk.Checkbutton(self.listbox_frame, text=os.path.basename(d), variable=var, command=self.check_limit)
            chk.pack(anchor="w")
            self.check_vars[d] = var
            
        self.btn_generate.config(state=tk.NORMAL)
        
    def check_limit(self):
        selected = [d for d, var in self.check_vars.items() if var.get()]
        if len(selected) > 5:
            messagebox.showwarning("Limit Reached", "You can select up to 5 replicates.")
            # Uncheck the last one? Hard to know which was last, just warn
            
    def generate(self):
        selected = [d for d, var in self.check_vars.items() if var.get()]
        if not selected:
            messagebox.showwarning("Select Replicates", "Please select at least one replicate.")
            return
        if len(selected) > 5:
            messagebox.showwarning("Limit Reached", "You can only select up to 5 replicates.")
            return
            
        process_tests(self.root_folder, selected)

if __name__ == "__main__":
    root = tk.Tk()
    app = App(root)
    root.mainloop()
